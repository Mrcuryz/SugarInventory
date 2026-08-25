from pathlib import Path
import sys
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P


def iter_blocks(parent):
    parent_element = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


path = Path(r"D:\Users\Mrcury\Desktop\结合AI分析的水库智能无人巡检系统(1).docx")
sys.stdout.reconfigure(encoding="utf-8")
doc = Document(path)
print(f"FILE={path}")
print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)} INLINE_SHAPES={len(doc.inline_shapes)}")
for i, block in enumerate(iter_blocks(doc), start=1):
    if isinstance(block, Paragraph):
        text = " ".join(block.text.split())
        if text:
            print(f"P{i:04d}|STYLE={block.style.name}|{text}")
    else:
        print(f"TABLE{i:04d}|ROWS={len(block.rows)}|COLS={len(block.columns)}")
        for r_idx, row in enumerate(block.rows, start=1):
            cells = [" / ".join(" ".join(p.text.split()) for p in cell.paragraphs if p.text.strip()) for cell in row.cells]
            print(f"T{i:04d}R{r_idx:03d}|" + " || ".join(cells))

for s_idx, section in enumerate(doc.sections, start=1):
    header = " | ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    footer = " | ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    if header:
        print(f"HEADER{s_idx}|{header}")
    if footer:
        print(f"FOOTER{s_idx}|{footer}")
